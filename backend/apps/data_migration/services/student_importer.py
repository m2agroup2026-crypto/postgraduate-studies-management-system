import hashlib
import json
from datetime import date, datetime
from pathlib import Path

from django.db import transaction
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

from apps.academic_history.models import AcademicHistoryEvent
from apps.academics.models import AcademicDegree, Department, Program
from apps.committees.models import DefenseCommittee
from apps.data_migration.models import LegacyImportBatch, LegacyImportRow
from apps.students.models import AcademicEnrollment, Student
from apps.theses.models import Thesis

SOURCE_ID_PREFIX = "LEGACY-"
DEFAULT_SHEET = "Digitalization November  2022"


def clean(value):
    if value is None:
        return None
    if isinstance(value, str):
        value = " ".join(value.split())
        return value or None
    return value


def json_value(value):
    value = clean(value)
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return value


def identifier(value):
    value = clean(value)
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def parse_date(value):
    value = clean(value)
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        try:
            parsed = from_excel(value)
            return parsed.date() if isinstance(parsed, datetime) else parsed
        except (TypeError, ValueError, OverflowError):
            return None
    text = str(value).strip().replace(".", "/").replace("-", "/")
    for pattern in ("%d/%m/%Y", "%Y/%m/%d", "%d/%m/%y", "%Y"):
        try:
            parsed = datetime.strptime(text, pattern).date()
            return parsed
        except ValueError:
            continue
    return None


def truthy(value):
    if isinstance(value, bool):
        return value
    return str(clean(value) or "").casefold() in {"1", "true", "yes", "نعم", "صح"}


def stable_code(prefix, value, length=12):
    digest = hashlib.sha1(str(value).encode("utf-8"), usedforsecurity=False).hexdigest()[:length]
    return f"{prefix}-{digest}".upper()


def degree_definition(value):
    text = str(clean(value) or "").casefold()
    if "دكتور" in text or "phd" in text or "doctor" in text:
        return "PHD", "دكتوراه", "Doctorate", AcademicDegree.DegreeLevel.PHD
    if "دبلوم" in text or "diploma" in text:
        return "DIPLOMA", "دبلوم", "Diploma", AcademicDegree.DegreeLevel.DIPLOMA
    return "MASTER", "ماجستير", "Master", AcademicDegree.DegreeLevel.MASTER


def derived_thesis_status(row):
    if clean(row.get("تاريخ المنح")):
        return "COMPLETED"
    if clean(row.get("تاريخ المناقشة")):
        return "FINAL_APPROVED"
    return "REGISTERED"


STATUS_RANK = {
    "REGISTERED": 10,
    "SUBMITTED": 20,
    "UNDER_REVIEW": 30,
    "DIRECTOR_APPROVED": 40,
    "VICE_DEAN_APPROVED": 50,
    "DEAN_APPROVED": 60,
    "FINAL_APPROVED": 70,
    "COMPLETED": 80,
}


class StudentWorkbookImporter:
    def __init__(self, source, sheet_name=None):
        self.source = Path(source).resolve()
        self.sheet_name = sheet_name or DEFAULT_SHEET

    def checksum(self):
        digest = hashlib.sha256()
        with self.source.open("rb") as handle:
            for chunk in iter(lambda: handle.read(1024 * 1024), b""):
                digest.update(chunk)
        return digest.hexdigest()

    def read_rows(self):
        workbook = load_workbook(self.source, read_only=True, data_only=True)
        if self.sheet_name not in workbook.sheetnames:
            workbook.close()
            raise ValueError(f"Worksheet not found: {self.sheet_name}")
        sheet = workbook[self.sheet_name]
        values = sheet.iter_rows(values_only=True)
        headers = [str(clean(value) or "") for value in next(values)]
        rows = []
        for row_number, values_row in enumerate(values, start=2):
            payload = {
                header: json_value(value)
                for header, value in zip(headers, values_row, strict=False)
                if header
            }
            if any(value is not None for value in payload.values()):
                rows.append((row_number, payload))
        workbook.close()
        return rows

    def validate(self, rows):
        seen = set()
        prepared = []
        for row_number, payload in rows:
            source_id = identifier(payload.get("ID"))
            errors = []
            if not source_id:
                errors.append("missing_source_id")
            elif source_id in seen:
                errors.append("duplicate_source_id")
            else:
                seen.add(source_id)
            if not clean(payload.get("الاسم")):
                errors.append("missing_student_name")
            if not clean(payload.get("القسم")) and not clean(payload.get("Department")):
                errors.append("missing_department")
            serialized = json.dumps(payload, ensure_ascii=False, sort_keys=True, default=str)
            prepared.append(
                {
                    "row_number": row_number,
                    "source_id": source_id or f"ROW-{row_number}",
                    "payload": payload,
                    "checksum": hashlib.sha256(serialized.encode("utf-8")).hexdigest(),
                    "errors": errors,
                }
            )
        return prepared

    def run(self, commit=False):
        prepared = self.validate(self.read_rows())
        valid = [item for item in prepared if not item["errors"]]
        existing_ids = set(
            Student.objects.filter(
                university_id__in=[f"{SOURCE_ID_PREFIX}{item['source_id']}" for item in valid]
            ).values_list("university_id", flat=True)
        )
        result = {
            "source_file": self.source.name,
            "source_checksum": self.checksum(),
            "sheet": self.sheet_name,
            "mode": "commit" if commit else "dry-run",
            "read": len(prepared),
            "valid": len(valid),
            "created": sum(
                f"{SOURCE_ID_PREFIX}{item['source_id']}" not in existing_ids for item in valid
            ),
            "updated": sum(
                f"{SOURCE_ID_PREFIX}{item['source_id']}" in existing_ids for item in valid
            ),
            "failed": len(prepared) - len(valid),
            "errors": self.error_summary(prepared),
        }
        if not commit:
            return result

        with transaction.atomic():
            batch, _ = LegacyImportBatch.objects.get_or_create(
                source_checksum=result["source_checksum"],
                defaults={
                    "source_file": self.source.name,
                    "status": LegacyImportBatch.Status.VALIDATING,
                },
            )
            batch.source_file = self.source.name
            batch.status = LegacyImportBatch.Status.VALIDATING
            batch.total_records = result["read"]
            batch.processed_records = 0
            batch.failed_records = 0
            batch.analysis_result = result
            batch.save()

            committed_created = 0
            committed_updated = 0
            for item in prepared:
                archive, _ = LegacyImportRow.objects.update_or_create(
                    batch=batch,
                    source_identifier=item["source_id"],
                    defaults={
                        "sheet_name": self.sheet_name,
                        "source_row_number": item["row_number"],
                        "row_checksum": item["checksum"],
                        "raw_payload": item["payload"],
                        "status": (
                            LegacyImportRow.Status.FAILED
                            if item["errors"]
                            else LegacyImportRow.Status.VALID
                        ),
                        "errors": item["errors"],
                        "student": None,
                    },
                )
                if item["errors"]:
                    continue
                try:
                    with transaction.atomic():
                        student, created = self.import_row(
                            item["payload"], item["source_id"], archive
                        )
                        archive.student = student
                        archive.status = (
                            LegacyImportRow.Status.IMPORTED
                            if created
                            else LegacyImportRow.Status.UPDATED
                        )
                        archive.errors = []
                        archive.save(update_fields=["student", "status", "errors", "updated_at"])
                        committed_created += int(created)
                        committed_updated += int(not created)
                except Exception as exc:  # preserve the row and continue the audited batch
                    archive.status = LegacyImportRow.Status.FAILED
                    archive.errors = [f"import_error:{type(exc).__name__}"]
                    archive.save(update_fields=["status", "errors", "updated_at"])

            result["created"] = committed_created
            result["updated"] = committed_updated
            result["failed"] = batch.rows.filter(status=LegacyImportRow.Status.FAILED).count()
            batch.processed_records = committed_created + committed_updated
            batch.failed_records = result["failed"]
            batch.status = (
                LegacyImportBatch.Status.IMPORTED
                if batch.processed_records
                else LegacyImportBatch.Status.FAILED
            )
            batch.analysis_result = result
            batch.save(
                update_fields=[
                    "source_file",
                    "status",
                    "total_records",
                    "processed_records",
                    "failed_records",
                    "analysis_result",
                    "updated_at",
                ]
            )
        result["batch_id"] = batch.id
        return result

    @staticmethod
    def error_summary(prepared):
        summary = {}
        for item in prepared:
            for error in item["errors"]:
                summary[error] = summary.get(error, 0) + 1
        return summary

    def import_row(self, row, source_id, archive):
        department_ar = clean(row.get("القسم")) or clean(row.get("Department"))
        department_en = clean(row.get("Department")) or ""
        department, _ = Department.objects.update_or_create(
            code=stable_code("LEG", department_ar),
            defaults={"name_ar": department_ar, "name_en": department_en, "is_active": True},
        )
        student, created = Student.objects.update_or_create(
            university_id=f"{SOURCE_ID_PREFIX}{source_id}",
            defaults={"name_ar": clean(row.get("الاسم")), "department": department},
        )

        degree_code, degree_ar, degree_en, degree_level = degree_definition(
            row.get("الدرجة") or row.get("Degree")
        )
        degree, _ = AcademicDegree.objects.update_or_create(
            code=degree_code,
            defaults={"name_ar": degree_ar, "name_en": degree_en, "level": degree_level},
        )
        program, _ = Program.objects.update_or_create(
            department=department,
            code=f"{department.code}-{degree.code}"[:30],
            defaults={
                "degree": degree,
                "name_ar": f"{degree_ar} - {department.name_ar}",
                "name_en": f"{degree_en} - {department.name_en or department.name_ar}",
                "is_active": True,
            },
        )
        grant_date = parse_date(row.get("تاريخ المنح"))
        cancelled = truthy(row.get("الغاء قيد؟")) or bool(clean(row.get("تاريخ الالغاء")))
        frozen = truthy(row.get("تجميد قيد؟")) and not clean(
            row.get("تاريخ فك التجميد(مجلس كلية اعتبارا من)")
        )
        enrollment_status = (
            "CANCELLED"
            if cancelled
            else "COMPLETED"
            if grant_date
            else "FROZEN"
            if frozen
            else "ACTIVE"
        )
        enrollment, _ = AcademicEnrollment.objects.update_or_create(
            student=student,
            program=program,
            defaults={
                "status": enrollment_status,
                "enrollment_date": parse_date(row.get("تاريخ القيد(مجلس كلية)")),
                "completion_date": grant_date,
                "notes": "Imported from audited legacy data",
            },
        )

        title = (
            clean(row.get("عنوان البحث بعد التغيير الجوهرى(عربى)"))
            or clean(row.get("عنوان البحث بعد التغيير(عربى)"))
            or clean(row.get("عنوان البحث باللغة العربية"))
        )
        thesis = None
        if title:
            derived_status = derived_thesis_status(row)
            try:
                thesis = student.thesis
                status = max(
                    (thesis.status, derived_status),
                    key=lambda value: STATUS_RANK.get(value, 0),
                )
                thesis.title_ar = title
                thesis.status = status
                thesis.save(update_fields=["title_ar", "status"])
            except Thesis.DoesNotExist:
                thesis = Thesis.objects.create(
                    student=student,
                    title_ar=title,
                    status=derived_status,
                )

        defense_date = parse_date(row.get("تاريخ المناقشة"))
        if thesis and defense_date:
            DefenseCommittee.objects.update_or_create(
                thesis=thesis,
                defaults={
                    "defense_date": defense_date,
                    "status": "COMPLETED" if grant_date else "SCHEDULED",
                },
            )

        self.sync_history(student, enrollment, thesis, defense_date, row, archive)
        return student, created

    @staticmethod
    def sync_history(student, enrollment, thesis, defense_date, row, archive):
        events = []
        if enrollment.enrollment_date:
            events.append(("ENROLLMENT_CREATED", "إنشاء القيد الأكاديمي"))
        if thesis:
            events.append(("THESIS_REGISTERED", "تسجيل الرسالة العلمية"))
        if defense_date:
            events.append(("DEFENSE_CREATED", "تشكيل لجنة وتحديد موعد المناقشة"))
        if clean(row.get("تاريخ المنح")):
            events.append(("DEFENSE_COMPLETED", "منح الدرجة العلمية"))
        milestones = [
            ("تجميد قيد؟", "تجميد القيد"),
            ("استمرار قيد؟", "استمرار القيد"),
            ("الغاء قيد؟", "إلغاء القيد"),
            ("اعادة قيد؟", "إعادة القيد"),
            ("تغيير جوهرى؟", "تغيير جوهري في الرسالة"),
            ("تغيير غير جوهرى", "تغيير غير جوهري في الرسالة"),
        ]
        for field, title in milestones:
            if truthy(row.get(field)):
                events.append(("OTHER", title))

        for event_type, title in events:
            AcademicHistoryEvent.objects.update_or_create(
                student=student,
                event_type=event_type,
                title_ar=title,
                reference_type="LegacyImportRow",
                reference_id=archive.id,
                defaults={"description_ar": "Imported from audited legacy source"},
            )
